import customtkinter as ctk
from tkinter import filedialog, ttk
import os
import difflib
import itertools
import pandas as pd
from docx import Document
import PyPDF2

# Daftar ekstensi file kode pemrograman
PROGRAMMING_EXTENSIONS = [
    '.py', '.java', '.js', '.c', '.cpp', '.h', '.hpp', '.html', '.css',
    '.php', '.sql', '.rb', '.go', '.rs', '.kt', '.swift'
]

class FileReader:
    """Kelas untuk membaca berbagai jenis file."""
    
    @staticmethod
    def read_file(filepath):
        ext = os.path.splitext(filepath)[1].lower()
        try:
            if ext == '.docx':
                return FileReader.read_docx(filepath)
            elif ext == '.pdf':
                return FileReader.read_pdf(filepath)
            else:
                return FileReader.read_text(filepath)
        except Exception as e:
            raise IOError(f"Error membaca file {filepath}: {e}")
    
    @staticmethod
    def read_text(filepath):
        with open(filepath, 'r', encoding='utf-8') as file:
            return file.read()
    
    @staticmethod
    def read_docx(filepath):
        doc = Document(filepath)
        return '\n'.join([para.text for para in doc.paragraphs])
    
    @staticmethod
    def read_pdf(filepath):
        with open(filepath, 'rb') as file:
            reader = PyPDF2.PdfReader(file)
            text = ''
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted
        return text

class PlagiarismCalculator:
    """Kelas untuk menghitung kesamaan dan pengurangan nilai plagiasi."""
    
    def __init__(self, threshold=80, max_reduction=20):
        self.threshold = threshold
        self.max_reduction = max_reduction
    
    @staticmethod
    def calculate_similarity(text1, text2):
        return difflib.SequenceMatcher(None, text1, text2).ratio()
    
    def calculate_reduction(self, plagiarism_percent):
        if plagiarism_percent <= self.threshold:
            return 0.0
        elif plagiarism_percent >= 100:
            return self.max_reduction
        else:
            return ((plagiarism_percent - self.threshold) / (100 - self.threshold)) * self.max_reduction

class PlagiarismCheckerApp:
    """Kelas utama aplikasi pengecek plagiasi."""
    
    def __init__(self):
        self.app = ctk.CTk()
        self.app.geometry("510x300")
        self.app.title("Pengecek Plagiasi")
        self.app.resizable(False, False)
        
        ctk.set_appearance_mode("Dark")
        ctk.set_default_color_theme("blue")
        
        self.plagiarism_calculator = PlagiarismCalculator()
        self.files_content = {}
        
        self._create_widgets()
    
    def _create_widgets(self):
        # Label utama
        main_label = ctk.CTkLabel(self.app, text="Pengecek Plagiasi", font=ctk.CTkFont(size=18, weight="bold"))
        main_label.pack(pady=5)
        
        # Frame untuk input
        frame = ctk.CTkFrame(self.app)
        frame.pack(pady=5, padx=10, fill="both", expand=True)
        
        frame.grid_columnconfigure(1, weight=1)
        frame.grid_columnconfigure(2, weight=0)
        
        # Input file
        file_label = ctk.CTkLabel(frame, text="Pilih File:")
        file_label.grid(row=0, column=0, padx=5, pady=5, sticky="w")
        self.file_entry = ctk.CTkEntry(frame, placeholder_text="Path to files")
        self.file_entry.grid(row=0, column=1, padx=5, pady=5, sticky="ew")
        file_button = ctk.CTkButton(frame, text="Browse", command=self.browse_files)
        file_button.grid(row=0, column=2, padx=5, pady=5, sticky="e")
        
        # Threshold persentase kesamaan
        threshold_label = ctk.CTkLabel(frame, text="Threshold Kesamaan (%)")
        threshold_label.grid(row=1, column=0, padx=5, pady=5, sticky="w")
        self.threshold_entry = ctk.CTkEntry(frame, placeholder_text="80")
        self.threshold_entry.grid(row=1, column=1, columnspan=2, padx=5, pady=5, sticky="ew")
        
        # Pengurangan maksimal nilai
        max_reduction_label = ctk.CTkLabel(frame, text="Pengurangan Maksimal Nilai:")
        max_reduction_label.grid(row=2, column=0, padx=5, pady=5, sticky="w")
        self.max_reduction_entry = ctk.CTkEntry(frame, placeholder_text="20")
        self.max_reduction_entry.grid(row=2, column=1, columnspan=2, padx=5, pady=5, sticky="ew")
        
        # Pilih lokasi output
        output_label = ctk.CTkLabel(frame, text="Pilih Lokasi Output:")
        output_label.grid(row=3, column=0, padx=5, pady=5, sticky="w")
        self.output_entry = ctk.CTkEntry(frame, placeholder_text="Lokasi Output")
        self.output_entry.grid(row=3, column=1, padx=5, pady=5, sticky="ew")
        output_button = ctk.CTkButton(frame, text="Browse", command=self.browse_output_location)
        output_button.grid(row=3, column=2, padx=5, pady=5, sticky="e")
        
        # Tombol untuk memulai pengecekan plagiasi
        process_button = ctk.CTkButton(self.app, text="Mulai Pengecekan", command=self.start_process)
        process_button.pack(pady=10)
        
        # Kotak keterangan hasil
        result_frame = ctk.CTkFrame(self.app, fg_color="#2B2B2B")
        result_frame.pack(fill="x", padx=10, pady=5)
        
        self.result_label = ctk.CTkLabel(result_frame, text="", wraplength=500)
        self.result_label.pack(padx=5, pady=5)
    
    def browse_files(self):
        file_paths = filedialog.askopenfilenames(
            title="Pilih 2 atau lebih file", 
            filetypes=[("Supported files", 
                        "*.txt;*.docx;*.pdf;*.py;*.java;*.c;*.cpp;*.h;*.hpp;*.cs;*.js;*.html;*.css;*.php;*.sql;*.rb;*.pl;*.sh;*.swift;*.kt;*.go;*.rs;*.lua;*.r;*.m;*.json;*.xml;*.yaml;*.yml;*.ini;*.cfg;*.conf;*.md;*.rst;*.csv;*.tsv;*.xls;*.xlsx;*.ppt;*.pptx;*.odt;*.ods;*.odp")])
        if len(file_paths) < 2:
            self.result_label.configure(text="Pastikan memilih setidaknya dua file.")
            return
        self.file_entry.delete(0, 'end')
        self.file_entry.insert(0, ', '.join(file_paths))
    
    def browse_output_location(self):
        output_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx", 
            filetypes=[("Excel files", "*.xlsx")], 
            title="Pilih lokasi dan nama file output")
        if output_path:
            self.output_entry.delete(0, 'end')
            self.output_entry.insert(0, output_path)
    
    def start_process(self):
        files = self.file_entry.get().split(', ')
        output_file = self.output_entry.get()
        
        try:
            threshold = float(self.threshold_entry.get())
            max_reduction = float(self.max_reduction_entry.get())
            self.plagiarism_calculator.threshold = threshold
            self.plagiarism_calculator.max_reduction = max_reduction
        except ValueError:
            self.result_label.configure(text="Threshold dan pengurangan maksimal harus berupa angka.")
            return
        
        if not files or len(files) < 2:
            self.result_label.configure(text="Silakan pilih minimal dua file untuk diperiksa.")
            return
        
        if not output_file:
            self.result_label.configure(text="Silakan pilih lokasi untuk menyimpan hasil.")
            return
        
        # Panggil proses plagiarisme
        try:
            df_similarity, df_reduction = self.process_plagiarism(files)
            self.save_output(df_similarity, df_reduction, output_file)
            self.show_output_window(df_similarity, df_reduction)
        except IOError as e:
            self.result_label.configure(text=str(e))
        except Exception as e:
            self.result_label.configure(text=f"Terjadi kesalahan: {e}")
    
    def process_plagiarism(self, files):
        if len(files) < 2:
            raise ValueError("Pastikan ada setidaknya dua file kode untuk dibandingkan.")
        
        self.files_content = {}
        reduction_dict = {os.path.basename(path): 0.0 for path in files}
        error_files = set()
        
        for path in files:
            try:
                content = FileReader.read_file(path)
                self.files_content[os.path.basename(path)] = content
            except Exception as e:
                reduction_dict[os.path.basename(path)] = "Error"
                error_files.add(path)
                self.result_label.configure(text=f"Error membaca file {path}: {e}")
        
        similarities = []
        file_pairs = itertools.combinations(self.files_content.keys(), 2)
        for file1, file2 in file_pairs:
            text1 = self.files_content[file1]
            text2 = self.files_content[file2]
            similarity_ratio = self.plagiarism_calculator.calculate_similarity(text1, text2) * 100
            similarities.append((file1, file2, round(similarity_ratio, 2)))
            
            if similarity_ratio > self.plagiarism_calculator.threshold:
                reduction = self.plagiarism_calculator.calculate_reduction(similarity_ratio)
                reduction_dict[file1] = max(reduction_dict[file1], reduction) if isinstance(reduction_dict[file1], float) else reduction_dict[file1]
                reduction_dict[file2] = max(reduction_dict[file2], reduction) if isinstance(reduction_dict[file2], float) else reduction_dict[file2]
        
        # Membuat DataFrame untuk pengurangan nilai
        data_reduction = [{
            "Nama File": student,
            "Persentase Pengurangan Nilai (%)": round(reduction, 2) if isinstance(reduction, float) else reduction
        } for student, reduction in reduction_dict.items()]
        df_reduction = pd.DataFrame(data_reduction).sort_values(by="Persentase Pengurangan Nilai (%)", ascending=False)
        
        # Membuat DataFrame untuk kesamaan
        data_similarity = [{
            "File 1": file1,
            "File 2": file2,
            "Kesamaan (%)": sim
        } for file1, file2, sim in similarities]
        df_similarity = pd.DataFrame(data_similarity).sort_values(by="Kesamaan (%)", ascending=False)
        
        self.df_similarity = df_similarity
        self.df_reduction = df_reduction
        
        return df_similarity, df_reduction
    
    def save_output(self, df_similarity, df_reduction, output_file):
        try:
            with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
                df_reduction.to_excel(writer, sheet_name='Pengurangan Nilai', index=False)
                df_similarity.to_excel(writer, sheet_name='Kesamaan', index=False)
            self.result_label.configure(text=f"Output berhasil disimpan ke {output_file}")
        except Exception as e:
            raise IOError(f"Error menyimpan file Excel: {e}")
    
    def show_output_window(self, df_similarity, df_reduction):
        if df_similarity is None or df_reduction is None:
            return
        
        output_window = ctk.CTkToplevel(self.app)
        output_window.title("Hasil Pengecekan Plagiasi")
        output_window.geometry("800x600")
        output_window.attributes("-topmost", True)
        
        # Tabel Kesamaan
        similarity_label = ctk.CTkLabel(output_window, text="Persentase Kesamaan", font=ctk.CTkFont(size=18, weight="bold"))
        similarity_label.pack(pady=5)
        
        similarity_frame = ctk.CTkFrame(output_window)
        similarity_frame.pack(pady=5, padx=10, fill="both", expand=True)
        
        similarity_scroll = ctk.CTkScrollbar(similarity_frame, orientation='vertical')
        similarity_scroll.pack(side="right", fill="y")
        
        similarity_columns = [str(col) for col in df_similarity.columns]
        similarity_table = ttk.Treeview(similarity_frame, columns=similarity_columns, show='headings', height=8, yscrollcommand=similarity_scroll.set)
        
        style = ttk.Style()
        style.theme_use("default")
        style.configure("Treeview",
                        background="#2a2d2e",
                        foreground="white",
                        rowheight=35,
                        fieldbackground="#343638",
                        bordercolor="#343638",
                        borderwidth=0,
                        font=('Segoe UI', 14))
        style.map('Treeview', background=[('selected', '#22559b')])
        
        style.configure("Treeview.Heading",
                        background="#565b5e",
                        foreground="white",
                        relief="flat",
                        font=('Segoe UI', 14))
        style.map("Treeview.Heading",
                    background=[('active', '#3484F0')])
        
        similarity_scroll.configure(command=similarity_table.yview)
        
        for col in similarity_columns:
            similarity_table.heading(col, text=col)
            similarity_table.column(col, anchor='center', width=150)
        
        for _, row in df_similarity.iterrows():
            similarity_table.insert('', 'end', values=list(row))
        
        similarity_table.pack(fill="both", expand=True)
        
        # Tabel Pengurangan Nilai
        reduction_label = ctk.CTkLabel(output_window, text="Pengurangan Nilai", font=ctk.CTkFont(size=18, weight="bold"))
        reduction_label.pack(pady=5)
        
        reduction_frame = ctk.CTkFrame(output_window)
        reduction_frame.pack(pady=5, padx=10, fill="both", expand=True)
        
        reduction_scroll = ctk.CTkScrollbar(reduction_frame, orientation='vertical')
        reduction_scroll.pack(side="right", fill="y")
        
        reduction_columns = [str(col) for col in df_reduction.columns]
        reduction_table = ttk.Treeview(reduction_frame, columns=reduction_columns, show='headings', height=8, yscrollcommand=reduction_scroll.set)
        
        reduction_scroll.configure(command=reduction_table.yview)
        
        for col in reduction_columns:
            reduction_table.heading(col, text=col)
            reduction_table.column(col, anchor='center', width=150)
        
        for _, row in df_reduction.iterrows():
            reduction_table.insert('', 'end', values=list(row))
        
        reduction_table.pack(fill="both", expand=True)
        
        # Event double click untuk menampilkan perbandingan teks
        similarity_table.bind("<Double-1>", self.on_double_click_similarity)
    
    def on_double_click_similarity(self, event):
        selected_item = event.widget.selection()
        if not selected_item:
            return
        
        item = event.widget.item(selected_item)
        file1, file2 = item['values'][0], item['values'][1]
        
        text1 = self.files_content[file1]
        text2 = self.files_content[file2]
        
        self.show_comparison_window(file1, file2, text1, text2)
    
    def show_comparison_window(self, file1, file2, text1, text2):
        comparison_window = ctk.CTkToplevel(self.app)
        comparison_window.title(f"Perbandingan: {file1} vs {file2}")
        comparison_window.geometry("900x600")
        comparison_window.attributes("-topmost", True)
        
        # Label judul
        title_label = ctk.CTkLabel(comparison_window, text=f"Perbandingan: {file1} vs {file2}", font=ctk.CTkFont(size=18, weight="bold"))
        title_label.pack(pady=10)
        
        # Frame untuk teks
        text_frame = ctk.CTkFrame(comparison_window)
        text_frame.pack(fill="both", expand=True, padx=10, pady=10)
        
        # Frame kiri untuk file 1
        text1_frame = ctk.CTkFrame(text_frame)
        text1_frame.pack(side="left", fill="both", expand=True, padx=10, pady=10)
        
        # Text widget untuk menampilkan teks file 1
        text1_widget = ctk.CTkTextbox(text1_frame, wrap="none")
        text1_widget.pack(side="left", fill="both", expand=True)
        text1_widget.insert("1.0", text1)
        text1_widget.configure(state="disabled")
        
        self.set_programming_font(text1_widget, file1)
        
        # Frame kanan untuk file 2
        text2_frame = ctk.CTkFrame(text_frame)
        text2_frame.pack(side="right", fill="both", expand=True, padx=10, pady=10)
        
        # Text widget untuk menampilkan teks file 2
        text2_widget = ctk.CTkTextbox(text2_frame, wrap="none")
        text2_widget.pack(side="left", fill="both", expand=True)
        text2_widget.insert("1.0", text2)
        text2_widget.configure(state="disabled")
        
        self.set_programming_font(text2_widget, file2)
        
        # Scrollbar sinkron untuk kedua teks
        sync_scroll_verticalbar = ctk.CTkScrollbar(text_frame, command=lambda *args: self.sync_scroll_vertical(text1_widget, text2_widget, *args))
        sync_scroll_verticalbar.pack(side="right", fill="y")
        
        text1_widget.configure(yscrollcommand=lambda *args: sync_scroll_verticalbar.set(*args))
        text2_widget.configure(yscrollcommand=lambda *args: sync_scroll_verticalbar.set(*args))
        
        # Scrollbar horizontal sinkron
        sync_scrollbar_horizontal = ctk.CTkScrollbar(comparison_window, command=lambda *args: self.sync_scroll_horizontal(text1_widget, text2_widget, *args), orientation="horizontal")
        sync_scrollbar_horizontal.pack(side="bottom", fill="x")
        
        text1_widget.configure(xscrollcommand=lambda *args: sync_scrollbar_horizontal.set(*args))
        text2_widget.configure(xscrollcommand=lambda *args: sync_scrollbar_horizontal.set(*args))
        
        # Highlight bagian yang sama
        self.highlight_similarities(text1_widget, text2_widget, text1, text2)
    
    @staticmethod
    def set_programming_font(widget, filepath):
        ext = os.path.splitext(filepath)[1].lower()
        if ext in PROGRAMMING_EXTENSIONS:
            widget.configure(font=("Consolas", 12))
        else:
            widget.configure(font=("Open Sans", 12))
    
    @staticmethod
    def sync_scroll_vertical(text1_widget, text2_widget, *args):
        text1_widget.yview(*args)
        text2_widget.yview(*args)
    
    @staticmethod
    def sync_scroll_horizontal(text1_widget, text2_widget, *args):
        text1_widget.xview(*args)
        text2_widget.xview(*args)
    
    @staticmethod
    def highlight_similarities(text1_widget, text2_widget, text1, text2):
        sequence_matcher = difflib.SequenceMatcher(None, text1, text2)
        
        for match in sequence_matcher.get_matching_blocks():
            start1, start2, length = match
            if length > 0:
                # Highlight di teks 1
                text1_widget.tag_add("highlight", f"1.0+{start1}c", f"1.0+{start1+length}c")
                text1_widget.tag_config("highlight", background="#4B5632", foreground="white")
                
                # Highlight di teks 2
                text2_widget.tag_add("highlight", f"1.0+{start2}c", f"1.0+{start2+length}c")
                text2_widget.tag_config("highlight", background="#4B5632", foreground="white")
    
    def run(self):
        self.app.mainloop()

if __name__ == "__main__":
    app = PlagiarismCheckerApp()
    app.run()
